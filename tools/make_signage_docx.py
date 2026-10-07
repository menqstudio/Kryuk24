"""Инструкция для типографии к комплекту вывесок (brand/06_signage) — Word.
Запуск системным Python с python-docx: py -3.12 tools/make_signage_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

SIG = Path(__file__).resolve().parents[1] / "brand" / "06_signage"
OUT = SIG / "КРЮК24_вывески_инструкция_для_печати.docx"

ORANGE = RGBColor(0xEF, 0x5B, 0x00)
INK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(20)
sec.top_margin = sec.bottom_margin = Mm(18)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(6)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def h(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = ORANGE
    return p


def para(parts, size=11, color=None, after=6):
    """parts: list of (text, bold) tuples or a plain string."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if isinstance(parts, str):
        parts = [(parts, False)]
    for t, b in parts:
        r = p.add_run(t)
        r.bold = b
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def bullet(parts):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if isinstance(parts, str):
        parts = [(parts, False)]
    for t, b in parts:
        r = p.add_run(t)
        r.bold = b
    return p


# ---- title
p = doc.add_paragraph()
r = p.add_run("КРЮК24")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = INK
p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph()
r = p.add_run("Вывески и таблички для помещения — что печатать и куда вешать")
r.bold = True
r.font.size = Pt(15)
r.font.color.rgb = ORANGE
p.paragraph_format.space_after = Pt(10)

para("Комплект нужен, чтобы оформить место приема клиентов и снять видео, по которому "
     "Яндекс Бизнес подтверждает адрес компании.")
para([("Срок подтверждения адреса в Яндексе — до 13.10.2026.", True)])

# ---- what to give the print shop
h("Что отдавать в типографию")
para([("Папка «PDF для печати»", True), (" — эти файлы и нужно нести в типографию. Это вектор, размер страницы "
      "1:1 в миллиметрах, шрифты встроены в файл: устанавливать их типографии не нужно.", False)])
para([("Папка «SVG исходники»", True), (" — те же макеты в исходном векторе. Отдавать, если типография попросит "
      "или захочет поправить размер у себя.", False)])
para("Цвета: оранжевый #EF5B00, темно-синий #13233A, белый. Макеты без вылетов и меток реза — "
     "если типографии нужны вылеты, фон просто продлевается тем же цветом.")

# ---- list of layouts
h("Макеты")
rows = [
    ("1", "Фасадная вывеска", "01_vyveska_fasad_2000x500mm.pdf", "2000 × 500 мм",
     "Фасад, над входом или рядом"),
    ("2", "Табличка у входа", "02_tablichka_vhod_400x300mm.pdf", "400 × 300 мм",
     "На дверь или на стену у входа"),
    ("3", "Указатель со стрелкой", "03_ukazatel_600x200mm_vpravo.pdf\n03_ukazatel_600x200mm_vlevo.pdf",
     "600 × 200 мм", "На путь от адресной таблички дома до входа, если вход не виден с улицы"),
    ("4", "Логотип в зону приема", "04_logo_zona_priema_1200x380mm.pdf", "1200 × 380 мм",
     "На стену за столом, где принимают клиента"),
    ("5", "Стенд «Тарифы»", "05_stend_tarify_A3_297x420mm.pdf", "297 × 420 мм (A3)",
     "Информационный стенд внутри помещения"),
    ("6", "Настольная табличка", "06_nastolnaya_tablichka_A5_210x148mm.pdf", "210 × 148 мм (A5)",
     "На стол — место приема заявок и оплаты"),
]
widths = [Mm(9), Mm(34), Mm(62), Mm(27), Mm(38)]
t = doc.add_table(rows=1, cols=5)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
for i, name in enumerate(["№", "Что это", "Файл", "Размер", "Куда"]):
    c = t.rows[0].cells[i]
    c.width = widths[i]
    shade(c, "13233A")
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = c.paragraphs[0].add_run(name)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for row in rows:
    cells = t.add_row().cells
    for i, val in enumerate(row):
        cells[i].width = widths[i]
        lines = val.split("\n")
        for j, line in enumerate(lines):
            pp = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
            pp.paragraph_format.space_after = Pt(0)
            r = pp.add_run(line)
            r.font.size = Pt(8.5 if i == 2 else 10)
            r.bold = i == 1

# ---- before printing
h("Перед печатью")
bullet([("Размеры макетов 1, 3 и 4 — исходные, по месту не замерены.", True),
        (" Сначала измерить место над входом и стену внутри, потом назвать типографии итоговый "
         "размер. Вектор масштабируется без потери качества, главное — сохранить пропорции: "
         "макет 1 — 4:1, макет 3 — 3:1, макет 4 — 60:19.", False)])
bullet([("Стрелку указателя выбрать по факту:", True),
        (" файл «vpravo» — стрелка вправо, «vlevo» — влево. Печатать один, нужный.", False)])
bullet("Материал и крепление — на выбор типографии под место: улица или помещение.")

# ---- what is written on the layouts
h("Что написано на макетах")
bullet("Телефон +7 985 893-06-06, сайт kryuk24.ru, «круглосуточно» — как в карточке Яндекс Бизнеса.")
bullet("Тарифы на стенде — те же, что на сайте kryuk24.ru. Если тариф изменился — сообщить до печати, "
       "макет будет исправлен.")
bullet("«Оплата: наличными или банковским переводом» — как указано в карточке Яндекс Бизнеса.")
bullet("Адреса, ИНН и наименования ИП или самозанятого на макетах нет. Если нужно добавить — "
       "прислать точный текст.")

# ---- reviews
h("Отзывы")
para("QR-наклейка и визитка для сбора отзывов уже готовы — это отдельные файлы. "
     "Наклейку 100 × 100 мм можно поставить на стол рядом с настольной табличкой (макет 6).")

# ---- previews
doc.add_page_break()
h("Как выглядят макеты")
previews = [
    ("1. Фасадная вывеска — 2000 × 500 мм", "01_vyveska_fasad_2000x500mm.png", 170),
    ("2. Табличка у входа — 400 × 300 мм", "02_tablichka_vhod_400x300mm.png", 80),
    ("3. Указатель — 600 × 200 мм (вариант «вправо»)", "03_ukazatel_600x200mm_vpravo.png", 120),
    ("3. Указатель — 600 × 200 мм (вариант «влево»)", "03_ukazatel_600x200mm_vlevo.png", 120),
    ("4. Логотип в зону приема — 1200 × 380 мм", "04_logo_zona_priema_1200x380mm.png", 150),
    ("5. Стенд «Тарифы» — A3", "05_stend_tarify_A3_297x420mm.png", 95),
    ("6. Настольная табличка — A5", "06_nastolnaya_tablichka_A5_210x148mm.png", 100),
]
for cap, f, w in previews:
    p = para([(cap, True)], size=10, color=GREY, after=3)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pic.add_run().add_picture(str(SIG / "png" / f), width=Mm(w))

doc.save(OUT)
print("saved", OUT.stat().st_size // 1024, "KB")
